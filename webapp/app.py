from __future__ import annotations

import io
import os
from datetime import datetime, timedelta
from functools import wraps

from flask import (
    Flask,
    flash,
    g,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from werkzeug.middleware.proxy_fix import ProxyFix

from webapp import db
from webapp.i18n import tr as i18n_tr
from webapp.modules_config import MODULES, SECTION_META, modules_for_section
from webapp import review_engine
from webapp import permissions
from webapp import warehouse_excel
from webapp import backup as backup_svc

app = Flask(__name__, instance_relative_config=True)
# يجب أن يبقى SECRET_KEY ثابتاً على Render — تغييره يُبطل جلسات الجميع
app.secret_key = os.environ.get("SECRET_KEY", "rakaz-khurais-emergency-2026")
app.config["TEMPLATES_AUTO_RELOAD"] = True
app.permanent_session_lifetime = timedelta(days=30)
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
# Render يضبط RENDER=true ويعمل خلف HTTPS
app.config["SESSION_COOKIE_SECURE"] = bool(os.environ.get("RENDER"))
app.url_map.strict_slashes = False
# Render / reverse proxies
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)

_DB_READY = False
PUBLIC_ENDPOINTS = {
    "login",
    "forgot_password",
    "set_lang",
    "static",
    "health",
    "api_backups_latest",
    "api_backups_auto_run",
    "api_backups_sync_status",
}


def create_app():
    global _DB_READY
    db.init_db()
    _DB_READY = True
    # ابدأ الحفظ التلقائي في الخلفية
    try:
        backup_svc.start_auto_backup_scheduler(app)
    except Exception:
        pass
    return app


@app.before_request
def _load_context():
    global _DB_READY
    if not _DB_READY:
        db.init_db()
        _DB_READY = True
    else:
        # بعد استعادة حفظة قديمة قد تختفي جداول جديدة
        try:
            db.ensure_schema()
        except Exception:
            pass
    g.settings = db.get_settings()
    g.lists = db.get_lists()
    g.year = datetime.now().year
    g.lang = session.get("lang") or "ar"
    # حماية الصفحات بنفس أسلوب أنظمة التقارير المشابهة
    if request.endpoint and request.endpoint not in PUBLIC_ENDPOINTS and not session.get("user_id"):
        if request.endpoint != "static":
            return redirect(url_for("login", next=request.path))
        return None
    # نظام الصلاحيات لكل التطبيق
    if session.get("user_id") and request.endpoint not in PUBLIC_ENDPOINTS:
        session["role"] = permissions.normalize_role(session.get("role"))
        missing = permissions.required_perm_for_request()
        if missing:
            label = permissions.PERM_LABELS.get(missing, missing)
            return permissions.deny_redirect(f"ليس لديك صلاحية: {label}")


def current_user_name():
    return session.get("full_name") or session.get("username") or "مستخدم"


def _after_data_change():
    """مزامنة صامتة بعد أي تعديل — بدون أزرار أو رسائل للمستخدم."""
    try:
        backup_svc.silent_backup_after_change()
    except Exception:
        pass


def login_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not session.get("user_id"):
            return redirect(url_for("login"))
        return fn(*args, **kwargs)

    return wrapper


@app.context_processor
def inject_globals():
    lang = session.get("lang") or "ar"

    def tr(key, **kwargs):
        return i18n_tr(lang, key, **kwargs)

    def can(*perms):
        return permissions.can(*perms)

    return {
        "settings": g.get("settings") or db.get_settings(),
        "lists": g.get("lists") or db.get_lists(),
        "current_year": g.get("year") or datetime.now().year,
        "logo_rekaz": url_for("static", filename="brand/rekaz.png"),
        "logo_rtc": url_for("static", filename="brand/rtc.jpg"),
        "logo_sec": url_for("static", filename="brand/sec.jpg"),
        "app_title": tr("app_title"),
        "lang": lang,
        "tr": tr,
        "can": can,
        "has_perm": permissions.has_perm,
        "nav_sections": permissions.nav_sections_for_role() if session.get("user_id") else [],
        "is_login_page": (request.endpoint or "") in {"login", "forgot_password"},
        "hosting": backup_svc.hosting_info(),
    }


def money(n):
    try:
        return f"{float(n or 0):,.2f} ر.س"
    except Exception:
        return "0.00 ر.س"


def response_minutes(dispatch, arrival):
    if not dispatch or not arrival:
        return None
    try:
        h1, m1 = map(int, dispatch.split(":")[:2])
        h2, m2 = map(int, arrival.split(":")[:2])
        diff = h2 * 60 + m2 - (h1 * 60 + m1)
        if diff < 0:
            diff += 24 * 60
        return diff
    except Exception:
        return None


def final_value(items_value, ratio=None):
    if items_value is None or items_value == "":
        return None
    ratio = ratio if ratio is not None else float(g.settings.get("emergency_ratio") or 0)
    return float(items_value) * (1 + ratio)


app.jinja_env.filters["money"] = money


def calc_cashflow(settings=None, cash_actual=None):
    settings = settings or g.settings
    monthly = (
        float(settings.get("teams_count") or 0)
        * float(settings.get("daily_tickets") or 0)
        * float(settings.get("work_days") or 0)
        * float(settings.get("target_avg") or 0)
    )
    expenses = float(settings.get("monthly_expenses") or 0)
    delay = int(settings.get("cash_delay_months") or 0)
    if cash_actual is None:
        conn = db.connect()
        cash_actual = {
            r["month_index"]: r["amount"]
            for r in conn.execute("SELECT month_index, amount FROM cash_actual").fetchall()
        }
        conn.close()
    rows = []
    cumulative = 0.0
    for i in range(12):
        raw = cash_actual.get(i)
        if raw is not None and raw != "":
            collection = float(raw)
        else:
            collection = 0.0 if i < delay else monthly
        net = collection - expenses
        cumulative += net
        rows.append(
            {
                "month": f"شهر {i + 1}",
                "index": i,
                "approved": monthly,
                "collection": collection,
                "expenses": expenses,
                "net": net,
                "cumulative": cumulative,
                "status": "يحتاج تمويل" if cumulative < 0 else "آمن",
                "raw": raw,
            }
        )
    return rows


def dashboard_stats():
    conn = db.connect()
    tickets = db.rows_to_dicts(conn.execute("SELECT * FROM tickets").fetchall())
    invoices = db.rows_to_dicts(conn.execute("SELECT * FROM invoices").fetchall())
    photos = db.rows_to_dicts(conn.execute("SELECT * FROM photos").fetchall())
    metering = db.rows_to_dicts(conn.execute("SELECT * FROM metering").fetchall())
    conn.close()

    target = float(g.settings.get("response_target") or 30)
    by_status = {s: 0 for s in g.lists.get("execution_status", [])}
    delayed = 0
    tickets_value = 0.0
    for t in tickets:
        by_status[t.get("status") or ""] = by_status.get(t.get("status") or "", 0) + 1
        mins = response_minutes(t.get("dispatch_time"), t.get("arrival_time"))
        if mins is not None and mins > target:
            delayed += 1
        fv = final_value(t.get("items_value"))
        if fv is not None:
            tickets_value += fv

    def photo_incomplete(p):
        keys = ["before_shot", "during_shot", "after_shot", "quantities_shot", "location_shot"]
        return not all(p.get(k) == "نعم" for k in keys)

    cash = calc_cashflow()
    liquidity = abs(min(0, min(r["cumulative"] for r in cash)))

    return {
        "total": len(tickets),
        "done": sum(1 for t in tickets if t.get("status") in ("منفذ", "مغلق")),
        "closed": sum(1 for t in tickets if t.get("status") == "مغلق"),
        "delayed": delayed,
        "photos_incomplete": sum(1 for p in photos if photo_incomplete(p)),
        "metering_approved": sum(1 for m in metering if m.get("status") == "معتمد"),
        "sap_raised": sum(1 for i in invoices if i.get("sap_status") in ("مرفوع", "مقبول")),
        "tickets_value": tickets_value,
        "invoices_total": sum(float(i.get("value") or 0) for i in invoices),
        "collected": sum(float(i.get("collected") or 0) for i in invoices),
        "remaining": sum(float(i.get("value") or 0) - float(i.get("collected") or 0) for i in invoices),
        "liquidity": liquidity,
        "by_status": by_status,
        "cash": cash,
    }


# ---------- Auth ----------
@app.route("/set-lang/<lang>")
def set_lang(lang):
    if lang not in ("ar", "en"):
        lang = "ar"
    session["lang"] = lang
    return redirect(request.referrer or url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if session.get("user_id"):
        return redirect(url_for("ops_home"))
    if request.method == "POST":
        username = (request.form.get("username") or "").strip()
        password = (request.form.get("password") or "").strip()
        conn = db.connect()
        user = conn.execute(
            "SELECT * FROM users WHERE lower(username)=lower(?)",
            (username,),
        ).fetchone()
        conn.close()
        if not user or (user["password"] or "") != password:
            flash(i18n_tr(session.get("lang") or "ar", "bad_login"), "danger")
            return render_template("login.html")
        if not user["active"]:
            flash(i18n_tr(session.get("lang") or "ar", "inactive_user"), "danger")
            return render_template("login.html")
        session.clear()
        # جلسة دائمة دائماً (30 يوماً) حتى لا يخرج المستخدم عند إعادة نشر Render
        session.permanent = True
        session["user_id"] = user["id"]
        session["username"] = user["username"]
        session["full_name"] = user["full_name"]
        session["role"] = permissions.normalize_role(user["role"])
        session["lang"] = session.get("lang") or "ar"
        db.log_audit(user["full_name"], "دخول", "نظام", user["id"], user["username"])
        nxt = request.args.get("next") or url_for("ops_home")
        if not str(nxt).startswith("/") or nxt in {"/", "/login"}:
            nxt = url_for("ops_home")
        return redirect(nxt)
    return render_template("login.html")


@app.route("/forgot-password")
def forgot_password():
    flash(i18n_tr(session.get("lang") or "ar", "forgot_hint"), "ok")
    return redirect(url_for("login"))


@app.route("/logout")
def logout():
    if session.get("user_id"):
        db.log_audit(current_user_name(), "خروج", "نظام", session.get("user_id"))
    lang = session.get("lang")
    session.clear()
    if lang:
        session["lang"] = lang
    return redirect(url_for("login"))


# ---------- Main hubs (نفس تبويبات تقارير رسملة) ----------
def _count(table):
    conn = db.connect()
    try:
        n = conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
        return n
    except Exception:
        try:
            db.ensure_schema(conn)
            n = conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
            return n
        except Exception:
            return 0
    finally:
        conn.close()


def section_links(section):
    links = []
    for key, mod in modules_for_section(section):
        links.append(
            {
                "label": mod["title"],
                "href": url_for("module_list", name=key),
                "count": _count(mod["table"]),
                "key": key,
            }
        )
    return links


@app.route("/health")
def health():
    """فحص نبض للإبقاء على الخدمة مستيقظة على Render."""
    auto = {}
    try:
        st = backup_svc.auto_status()
        auto = {
            "enabled": st.get("enabled"),
            "last_backup_at": st.get("last_backup_at"),
            "next_due_minutes": st.get("next_due_minutes"),
        }
    except Exception:
        auto = {"enabled": backup_svc.auto_backup_enabled()}
    return {
        "ok": True,
        "app": "rekaz",
        "build": os.environ.get("RENDER_GIT_COMMIT")
        or os.environ.get("RAKAZ_BUILD")
        or "local",
        "features": {
            "review_removed": True,
            "rasmalah_tabs": True,
            "backups": True,
            "auto_backup": True,
            "s3_backup": backup_svc.s3_configured(),
            "trial_mode": os.environ.get("TRIAL_MODE", "").strip() in {"1", "true", "yes"},
        },
        "auto_backup": auto,
    }, 200


@app.route("/")
@login_required
def dashboard():
    return redirect(url_for("ops_home"))


@app.route("/ops")
@login_required
def ops_home():
    links = [
        {"label": "بلاغات الأعمال", "href": url_for("tickets_list"), "count": _count("tickets")},
        {"label": "الكميات", "href": url_for("module_list", name="quantities"), "count": _count("quantities")},
        {"label": "قائمة الصور", "href": url_for("module_list", name="photos"), "count": _count("photos")},
        {"label": "التمتير", "href": url_for("module_list", name="metering"), "count": _count("metering")},
        {"label": "فرق المهام العاجلة", "href": url_for("teams_page"), "count": _count("teams")},
        {"label": "إجراءات العمل (SOP)", "href": url_for("sop_page")},
        {"label": "القوائم المرجعية", "href": url_for("lists_page")},
    ]
    return render_template(
        "ops_home.html",
        stats=dashboard_stats(),
        links=links,
        total_count=_count("tickets"),
    )


@app.route("/constructions")
@login_required
def constructions_home():
    links = section_links("constructions")
    return render_template(
        "section_hub.html",
        title="الإنشاءات - التنفيذ",
        subtitle="متابعة معاملات الإنشاءات والتنفيذ وربطها ببلاغات المكتب.",
        links=links,
        section="constructions",
        section_modules=modules_for_section("constructions"),
        section_meta=SECTION_META["constructions"],
        total_count=sum(i.get("count") or 0 for i in links),
    )


@app.route("/contractors")
@login_required
def contractors_home():
    links = section_links("contractors")
    return render_template(
        "section_hub.html",
        title="المقاولين",
        subtitle="متابعة أعمال المقاولين وربطها ببلاغات المكتب — بنفس تبويب تقارير رسملة.",
        links=links,
        section="contractors",
        section_modules=modules_for_section("contractors"),
        section_meta=SECTION_META["contractors"],
        total_count=sum(i.get("count") or 0 for i in links),
    )


@app.route("/quality")
@login_required
def quality_home():
    links = section_links("quality")
    return render_template(
        "section_hub.html",
        title="التنسيقات والجودة",
        subtitle="التنسيقات الفنية وإخلاءات الأسفلت وفحوصات الجودة.",
        links=links,
        section="quality",
        section_modules=modules_for_section("quality"),
        section_meta=SECTION_META["quality"],
        total_count=sum(i.get("count") or 0 for i in links),
    )


@app.route("/safety")
@login_required
def safety_home():
    links = section_links("safety")
    return render_template(
        "section_hub.html",
        title="السلامة",
        subtitle="تصاريح العمل وبلاغات السلامة المرتبطة بالمواقع.",
        links=links,
        section="safety",
        section_modules=modules_for_section("safety"),
        section_meta=SECTION_META["safety"],
        total_count=sum(i.get("count") or 0 for i in links),
    )


@app.route("/warehouses")
@login_required
def warehouses_home():
    links = section_links("warehouses")
    links.append({"label": "أرصدة المواد", "href": url_for("warehouse_balances"), "count": _count("warehouse_items")})
    return render_template(
        "section_hub.html",
        title="المستودعات",
        subtitle="الأصناف ومعاملات المستودع وأرصدة المواد — بنفس فكرة تقارير رسملة.",
        links=links,
        section="warehouses",
        section_modules=modules_for_section("warehouses"),
        section_meta=SECTION_META["warehouses"],
        total_count=_count("warehouse_tx"),
    )


@app.route("/external-purchases")
@login_required
def external_purchases_home():
    links = section_links("external")
    return render_template(
        "section_hub.html",
        title="المشتريات الخارجية والعهد",
        subtitle="طلبات الشراء الخارجي ومتابعة العهد المسلمة للموظفين.",
        links=links,
        section="external",
        section_modules=modules_for_section("external"),
        section_meta=SECTION_META["external"],
        total_count=sum(i.get("count") or 0 for i in links),
    )


@app.route("/financial")
@login_required
def financial_home():
    links = section_links("financial")
    links.append({"label": "التدفق النقدي", "href": url_for("cashflow")})
    stats = dashboard_stats()
    finance_stats = {
        "sap_raised": stats["sap_raised"],
        "invoices_total": stats["invoices_total"],
        "collected": stats["collected"],
        "remaining": stats["remaining"],
        "liquidity": stats["liquidity"],
    }
    return render_template(
        "section_hub.html",
        title="المتابعات المالية",
        subtitle="المستخلصات و SAP والتدفق النقدي — بنفس تبويب تقارير رسملة.",
        links=links,
        section="financial",
        section_modules=modules_for_section("financial"),
        section_meta=SECTION_META["financial"],
        total_count=_count("invoices"),
        finance_stats=finance_stats,
    )


@app.route("/maintenance")
@login_required
def maintenance_home():
    links = section_links("maintenance")
    return render_template(
        "section_hub.html",
        title="الورشة (سيارات - معدات)",
        subtitle="متابعة سيارات ومعدات الورش وربطها بالفرق الميدانية.",
        links=links,
        section="maintenance",
        section_modules=modules_for_section("maintenance"),
        section_meta=SECTION_META["maintenance"],
        total_count=sum(i.get("count") or 0 for i in links),
    )


@app.route("/hr")
@login_required
def hr_home():
    links = section_links("hr")
    return render_template(
        "section_hub.html",
        title="الموارد البشرية",
        subtitle="سجل الموظفين والأقسام وحالات الالتحاق — بنفس تبويب تقارير رسملة.",
        links=links,
        section="hr",
        section_modules=modules_for_section("hr"),
        section_meta=SECTION_META["hr"],
        total_count=sum(i.get("count") or 0 for i in links),
    )


@app.route("/contracts-admin")
@login_required
def contracts_admin_home():
    links = section_links("contracts")
    return render_template(
        "section_hub.html",
        title="إدارة العقود",
        subtitle="عقود المكتب وحالاتها وقيمها.",
        links=links,
        section="contracts",
        section_modules=modules_for_section("contracts"),
        section_meta=SECTION_META["contracts"],
        total_count=sum(i.get("count") or 0 for i in links),
    )


@app.route("/users")
@login_required
def users_home():
    return redirect(url_for("users_list"))


@app.route("/admin/audit-log")
@login_required
def audit_log_home():
    return redirect(url_for("audit_log_page"))


# ---------- Tickets (معاملات البلاغات) ----------
TICKET_FIELDS = [
    "ticket_no",
    "receive_date",
    "district",
    "receive_time",
    "agent",
    "station_no",
    "feeder_no",
    "location",
    "fault_type",
    "classification",
    "team",
    "dispatch_time",
    "arrival_time",
    "status",
    "execution_date",
    "photographed",
    "quantities_done",
    "asphalt_clearance",
    "metering_status",
    "consultant_approval",
    "invoice_status",
    "work_order",
    "invoice_no",
    "sap_status",
    "items_value",
    "notes",
]


def ticket_from_form():
    data = {f: (request.form.get(f) or "").strip() for f in TICKET_FIELDS}
    iv = data.get("items_value")
    data["items_value"] = float(iv) if iv not in ("", None) else None
    return data


@app.route("/tickets")
@login_required
def tickets_list():
    q = (request.args.get("q") or "").strip()
    status = (request.args.get("status") or "").strip()
    conn = db.connect()
    sql = "SELECT * FROM tickets WHERE 1=1"
    params = []
    if q:
        sql += " AND (ticket_no LIKE ? OR district LIKE ? OR fault_type LIKE ? OR team LIKE ? OR agent LIKE ?)"
        like = f"%{q}%"
        params.extend([like, like, like, like, like])
    if status:
        sql += " AND status = ?"
        params.append(status)
    sql += " ORDER BY id DESC"
    rows = db.rows_to_dicts(conn.execute(sql, params).fetchall())
    conn.close()
    for r in rows:
        r["response_min"] = response_minutes(r.get("dispatch_time"), r.get("arrival_time"))
        r["final_value"] = final_value(r.get("items_value"))
    return render_template("tickets_list.html", rows=rows, q=q, status=status)


@app.route("/tickets/new", methods=["GET", "POST"])
@login_required
def ticket_new():
    if request.method == "POST":
        data = ticket_from_form()
        if not data["ticket_no"]:
            flash("رقم البلاغ مطلوب", "danger")
            return render_template("ticket_form.html", row=data, mode="new")
        conn = db.connect()
        try:
            cols = ", ".join(TICKET_FIELDS)
            placeholders = ", ".join(["?"] * len(TICKET_FIELDS))
            cur = conn.execute(
                f"INSERT INTO tickets({cols}) VALUES ({placeholders})",
                [data[f] for f in TICKET_FIELDS],
            )
            conn.commit()
            db.log_audit(current_user_name(), "إضافة", "بلاغ", cur.lastrowid, data.get("ticket_no"))
            flash("تم إنشاء المعاملة بنجاح", "ok")
            _after_data_change()
            return redirect(url_for("tickets_list"))
        except Exception as exc:
            flash(f"تعذر الحفظ: {exc}", "danger")
        finally:
            conn.close()
    blank = {f: "" for f in TICKET_FIELDS}
    blank["receive_date"] = datetime.now().strftime("%Y-%m-%d")
    blank["status"] = "جديد"
    return render_template("ticket_form.html", row=blank, mode="new")


@app.route("/tickets/<int:ticket_id>")
@login_required
def ticket_view(ticket_id):
    conn = db.connect()
    row = conn.execute("SELECT * FROM tickets WHERE id=?", (ticket_id,)).fetchone()
    if not row:
        conn.close()
        flash("المعاملة غير موجودة", "danger")
        return redirect(url_for("tickets_list"))
    ticket = dict(row)
    tno = ticket["ticket_no"]
    related = {
        "quantities": db.rows_to_dicts(conn.execute("SELECT * FROM quantities WHERE ticket_no=?", (tno,)).fetchall()),
        "photos": db.rows_to_dicts(conn.execute("SELECT * FROM photos WHERE ticket_no=?", (tno,)).fetchall()),
        "coordination": db.rows_to_dicts(conn.execute("SELECT * FROM coordination WHERE ticket_no=?", (tno,)).fetchall()),
        "metering": db.rows_to_dicts(conn.execute("SELECT * FROM metering WHERE ticket_no=?", (tno,)).fetchall()),
    }
    conn.close()
    ticket["response_min"] = response_minutes(ticket.get("dispatch_time"), ticket.get("arrival_time"))
    ticket["final_value"] = final_value(ticket.get("items_value"))
    return render_template("ticket_view.html", ticket=ticket, related=related)


@app.route("/tickets/<int:ticket_id>/edit", methods=["GET", "POST"])
@login_required
def ticket_edit(ticket_id):
    conn = db.connect()
    row = conn.execute("SELECT * FROM tickets WHERE id=?", (ticket_id,)).fetchone()
    if not row:
        conn.close()
        flash("المعاملة غير موجودة", "danger")
        return redirect(url_for("tickets_list"))
    if request.method == "POST":
        data = ticket_from_form()
        sets = ", ".join([f"{f}=?" for f in TICKET_FIELDS])
        conn.execute(
            f"UPDATE tickets SET {sets}, updated_at=CURRENT_TIMESTAMP WHERE id=?",
            [data[f] for f in TICKET_FIELDS] + [ticket_id],
        )
        conn.commit()
        conn.close()
        db.log_audit(current_user_name(), "تعديل", "بلاغ", ticket_id, data.get("ticket_no"))
        flash("تم حفظ المعاملة", "ok")
        _after_data_change()
        return redirect(url_for("ticket_view", ticket_id=ticket_id))
    ticket = dict(row)
    conn.close()
    return render_template("ticket_form.html", row=ticket, mode="edit")


@app.route("/tickets/<int:ticket_id>/delete", methods=["POST"])
@login_required
def ticket_delete(ticket_id):
    conn = db.connect()
    row = conn.execute("SELECT ticket_no FROM tickets WHERE id=?", (ticket_id,)).fetchone()
    conn.execute("DELETE FROM tickets WHERE id=?", (ticket_id,))
    conn.commit()
    conn.close()
    db.log_audit(current_user_name(), "حذف", "بلاغ", ticket_id, row["ticket_no"] if row else "")
    flash("تم حذف المعاملة", "ok")
    _after_data_change()
    return redirect(url_for("tickets_list"))


@app.route("/tickets/<int:ticket_id>/print")
@login_required
def ticket_print(ticket_id):
    conn = db.connect()
    row = conn.execute("SELECT * FROM tickets WHERE id=?", (ticket_id,)).fetchone()
    if not row:
        conn.close()
        flash("المعاملة غير موجودة", "danger")
        return redirect(url_for("tickets_list"))
    ticket = dict(row)
    tno = ticket["ticket_no"]
    quantities = db.rows_to_dicts(conn.execute("SELECT * FROM quantities WHERE ticket_no=?", (tno,)).fetchall())
    photos = db.rows_to_dicts(conn.execute("SELECT * FROM photos WHERE ticket_no=?", (tno,)).fetchall())
    coordination = db.rows_to_dicts(conn.execute("SELECT * FROM coordination WHERE ticket_no=?", (tno,)).fetchall())
    metering = db.rows_to_dicts(conn.execute("SELECT * FROM metering WHERE ticket_no=?", (tno,)).fetchall())
    conn.close()
    ticket["response_min"] = response_minutes(ticket.get("dispatch_time"), ticket.get("arrival_time"))
    ticket["final_value"] = final_value(ticket.get("items_value"))
    for q in quantities:
        q["total"] = (float(q.get("qty") or 0) * float(q.get("unit_price") or 0))
    return render_template(
        "ticket_print.html",
        ticket=ticket,
        quantities=quantities,
        photos=photos,
        coordination=coordination,
        metering=metering,
        printed_at=datetime.now().strftime("%d-%m-%Y %H:%M:%S"),
    )


# ---------- Generic CRUD helpers ----------
# MODULES imported from webapp.modules_config


def _module_form_data(module):
    data = {}
    for key, _label, ftype in module["fields"]:
        val = (request.form.get(key) or "").strip()
        if ftype == "number":
            data[key] = float(val) if val != "" else None
        else:
            data[key] = val
    return data


@app.route("/module/<name>")
@login_required
def module_list(name):
    module = MODULES.get(name)
    if not module:
        flash("القسم غير موجود", "danger")
        return redirect(url_for("ops_home"))
    conn = db.connect()
    rows = db.rows_to_dicts(conn.execute(f"SELECT * FROM {module['table']} ORDER BY id DESC").fetchall())
    tickets = [r["ticket_no"] for r in conn.execute("SELECT ticket_no FROM tickets ORDER BY id DESC").fetchall()]
    conn.close()
    if name == "quantities":
        for r in rows:
            r["total"] = (float(r.get("qty") or 0) * float(r.get("unit_price") or 0))
    if name == "photos":
        for r in rows:
            keys = ["before_shot", "during_shot", "after_shot", "quantities_shot", "location_shot"]
            r["complete"] = "مكتمل" if all(r.get(k) == "نعم" for k in keys) else "ناقص"
    if name == "invoices":
        for r in rows:
            r["remaining"] = float(r.get("value") or 0) - float(r.get("collected") or 0)
    if name == "external_purchases":
        for r in rows:
            r["total"] = (float(r.get("qty") or 0) * float(r.get("unit_price") or 0))
    item_filter = (request.args.get("item_no") or "").strip()
    ticket_filter = (request.args.get("ticket_no") or "").strip()
    if name == "warehouse_items":
        for r in rows:
            r["balance"] = db.warehouse_balance(r.get("item_no"))
    if name == "warehouse_tx":
        if item_filter:
            rows = [r for r in rows if (r.get("item_no") or "").lower() == item_filter.lower()]
        if ticket_filter:
            rows = [r for r in rows if (r.get("ticket_no") or "") == ticket_filter]
    section = module.get("section")
    return render_template(
        "module_list.html",
        name=name,
        module=module,
        rows=rows,
        tickets=tickets,
        item_filter=item_filter,
        ticket_filter=ticket_filter,
        section=section,
        section_meta=SECTION_META.get(section),
        section_modules=modules_for_section(section) if section else [],
    )


@app.route("/module/<name>/new", methods=["GET", "POST"])
@login_required
def module_new(name):
    module = MODULES.get(name)
    if not module:
        return redirect(url_for("ops_home"))
    conn = db.connect()
    tickets = [r["ticket_no"] for r in conn.execute("SELECT ticket_no FROM tickets ORDER BY id DESC").fetchall()]
    prefill = {f[0]: "" for f in module["fields"]}
    if request.args.get("ticket_no") and "ticket_no" in prefill:
        prefill["ticket_no"] = request.args.get("ticket_no")
    if request.method == "POST":
        data = _module_form_data(module)
        if name == "warehouse_tx":
            data = db.enrich_warehouse_tx_from_item(data)
        keys = [f[0] for f in module["fields"]]
        cur = conn.execute(
            f"INSERT INTO {module['table']}({', '.join(keys)}) VALUES ({', '.join(['?']*len(keys))})",
            [data[k] for k in keys],
        )
        conn.commit()
        new_id = cur.lastrowid
        conn.close()
        db.log_audit(current_user_name(), "إضافة", module["title"], new_id, str(data)[:240])
        flash("تمت الإضافة", "ok")
        _after_data_change()
        return redirect(url_for("module_list", name=name))
    warehouse_items = db.list_warehouse_items() if name == "warehouse_tx" else []
    if request.args.get("item_no") and "item_no" in prefill:
        prefill["item_no"] = request.args.get("item_no")
        prefill = db.enrich_warehouse_tx_from_item(prefill)
    conn.close()
    section = module.get("section")
    return render_template(
        "module_form.html",
        name=name,
        module=module,
        row=prefill,
        tickets=tickets,
        warehouse_items=warehouse_items,
        mode="new",
        section=section,
        section_meta=SECTION_META.get(section),
        section_modules=modules_for_section(section) if section else [],
    )


@app.route("/module/<name>/<int:row_id>/edit", methods=["GET", "POST"])
@login_required
def module_edit(name, row_id):
    module = MODULES.get(name)
    if not module:
        return redirect(url_for("ops_home"))
    conn = db.connect()
    row = conn.execute(f"SELECT * FROM {module['table']} WHERE id=?", (row_id,)).fetchone()
    tickets = [r["ticket_no"] for r in conn.execute("SELECT ticket_no FROM tickets ORDER BY id DESC").fetchall()]
    if not row:
        conn.close()
        flash("السجل غير موجود", "danger")
        return redirect(url_for("module_list", name=name))
    if request.method == "POST":
        data = _module_form_data(module)
        if name == "warehouse_tx":
            data = db.enrich_warehouse_tx_from_item(data)
        keys = [f[0] for f in module["fields"]]
        sets = ", ".join([f"{k}=?" for k in keys])
        conn.execute(
            f"UPDATE {module['table']} SET {sets} WHERE id=?",
            [data[k] for k in keys] + [row_id],
        )
        conn.commit()
        conn.close()
        db.log_audit(current_user_name(), "تعديل", module["title"], row_id, str(data)[:240])
        flash("تم الحفظ", "ok")
        _after_data_change()
        return redirect(url_for("module_list", name=name))
    data = dict(row)
    warehouse_items = db.list_warehouse_items() if name == "warehouse_tx" else []
    conn.close()
    section = module.get("section")
    return render_template(
        "module_form.html",
        name=name,
        module=module,
        row=data,
        tickets=tickets,
        warehouse_items=warehouse_items,
        mode="edit",
        section=section,
        section_meta=SECTION_META.get(section),
        section_modules=modules_for_section(section) if section else [],
    )


@app.route("/module/<name>/<int:row_id>/delete", methods=["POST"])
@login_required
def module_delete(name, row_id):
    module = MODULES.get(name)
    if not module:
        return redirect(url_for("ops_home"))
    conn = db.connect()
    conn.execute(f"DELETE FROM {module['table']} WHERE id=?", (row_id,))
    conn.commit()
    conn.close()
    db.log_audit(current_user_name(), "حذف", module["title"], row_id)
    flash("تم الحذف", "ok")
    return redirect(url_for("module_list", name=name))


# ---------- Cashflow / Teams / Lists / Settings / SOP ----------
@app.route("/cashflow", methods=["GET", "POST"])
@login_required
def cashflow():
    conn = db.connect()
    if request.method == "POST":
        for i in range(12):
            raw = (request.form.get(f"m{i}") or "").strip()
            amount = float(raw) if raw != "" else None
            conn.execute(
                "INSERT INTO cash_actual(month_index, amount) VALUES (?,?) ON CONFLICT(month_index) DO UPDATE SET amount=excluded.amount",
                (i, amount),
            )
        conn.commit()
        flash("تم حفظ التحصيل الشهري", "ok")
        _after_data_change()
    cash_actual = {r["month_index"]: r["amount"] for r in conn.execute("SELECT * FROM cash_actual").fetchall()}
    conn.close()
    rows = calc_cashflow(cash_actual=cash_actual)
    return render_template(
        "cashflow.html",
        rows=rows,
        section="financial",
        section_meta=SECTION_META["financial"],
        section_modules=modules_for_section("financial"),
    )


@app.route("/teams", methods=["GET", "POST"])
@login_required
def teams_page():
    conn = db.connect()
    if request.method == "POST":
        action = request.form.get("action")
        if action == "add":
            conn.execute(
                "INSERT INTO teams(name,leader,technicians,driver,vehicle,area,status,notes) VALUES (?,?,?,?,?,?,?,?)",
                (
                    request.form.get("name"),
                    request.form.get("leader"),
                    request.form.get("technicians") or 0,
                    request.form.get("driver"),
                    request.form.get("vehicle"),
                    request.form.get("area"),
                    request.form.get("status") or "نشطة",
                    request.form.get("notes"),
                ),
            )
            conn.commit()
            flash("تمت إضافة الفرقة", "ok")
        elif action == "delete":
            conn.execute("DELETE FROM teams WHERE id=?", (request.form.get("id"),))
            conn.commit()
            flash("تم الحذف", "ok")
    rows = db.rows_to_dicts(conn.execute("SELECT * FROM teams ORDER BY id").fetchall())
    conn.close()
    return render_template("teams.html", rows=rows)


@app.route("/lists", methods=["GET", "POST"])
@login_required
def lists_page():
    if request.method == "POST":
        data = {}
        for key in db.DEFAULT_LISTS:
            text = request.form.get(key) or ""
            data[key] = [x.strip() for x in text.splitlines() if x.strip()]
        db.save_lists(data)
        g.lists = db.get_lists()
        flash("تم حفظ القوائم", "ok")
        _after_data_change()
    return render_template("lists.html", lists_data=db.get_lists())


@app.route("/settings", methods=["GET", "POST"])
@login_required
def settings_page():
    if request.method == "POST":
        data = {
            "office_name": request.form.get("office_name") or "",
            "company_name": request.form.get("company_name") or "",
            "teams_count": int(request.form.get("teams_count") or 0),
            "daily_tickets": int(request.form.get("daily_tickets") or 0),
            "work_days": int(request.form.get("work_days") or 0),
            "monthly_expenses": float(request.form.get("monthly_expenses") or 0),
            "cash_delay_months": int(request.form.get("cash_delay_months") or 0),
            "emergency_ratio": float(request.form.get("emergency_ratio") or 0),
            "target_avg": float(request.form.get("target_avg") or 0),
            "reject_limit": float(request.form.get("reject_limit") or 0),
            "response_target": float(request.form.get("response_target") or 0),
            "invoice_days": int(request.form.get("invoice_days") or 0),
        }
        db.save_settings(data)
        g.settings = db.get_settings()
        flash("تم حفظ الإعدادات", "ok")
        _after_data_change()
    return render_template("settings.html", s=db.get_settings())


@app.route("/admin/backups")
@login_required
def backups_home():
    backups = backup_svc.list_backups(limit=80)
    sizes = {}
    for b in backups:
        try:
            sizes[b["_rel"]] = backup_svc.human_size((b.get("progress") or {}).get("db_size_bytes") or 0)
        except Exception:
            sizes[b["_rel"]] = "—"
    current = backup_svc.progress_snapshot()
    auto = backup_svc.auto_status()
    sync_token = backup_svc.get_sync_token()
    base = request.url_root.rstrip("/")
    return render_template(
        "backups.html",
        backups=backups,
        timeline=backup_svc.progress_timeline(limit=40),
        purposes=backup_svc.PURPOSE_CHOICES,
        data_root=str(backup_svc.data_root()),
        backups_root=str(backup_svc.backups_root()),
        current_progress=current,
        current_size=backup_svc.human_size(current.get("db_size_bytes") or 0),
        sizes=sizes,
        hosting=backup_svc.hosting_info(),
        auto=auto,
        sync_token=sync_token,
        sync_latest_url=f"{base}/api/backups/latest?token={sync_token}",
        sync_status_url=f"{base}/api/backups/sync-status?token={sync_token}",
        sync_autorun_url=f"{base}/api/backups/auto-run?token={sync_token}",
    )


@app.route("/admin/backups/create", methods=["POST"])
@login_required
def backups_create():
    try:
        purpose = request.form.get("purpose") or "manual"
        if backup_svc.is_trial_free() and purpose == "manual":
            purpose = "trial"
        meta = backup_svc.create_backup(
            label=request.form.get("label") or "",
            note=request.form.get("note") or "",
            purpose=purpose,
            user_name=current_user_name(),
        )
        db.log_audit(
            current_user_name(),
            "حفظ بيانات",
            "نسخة احتياطية",
            details=f"{meta.get('purpose_label')} — {meta.get('label')} — بلاغات {meta.get('progress', {}).get('tickets_total', 0)}",
        )
        flash(f"تم إنشاء الحفظة: {meta.get('label') or meta.get('purpose_label')}", "ok")
    except Exception as exc:
        flash(f"تعذر إنشاء الحفظة: {exc}", "danger")
    return redirect(url_for("backups_home"))


@app.route("/admin/backups/download")
@login_required
def backups_download():
    rel = request.args.get("rel") or ""
    try:
        zip_path = backup_svc.build_backup_zip(rel)
        return send_file(
            zip_path,
            as_attachment=True,
            download_name=f"rekaz-backup-{rel.replace('/', '-')}.zip",
        )
    except Exception as exc:
        flash(f"تعذر التنزيل: {exc}", "danger")
        return redirect(url_for("backups_home"))


@app.route("/admin/backups/restore", methods=["POST"])
@login_required
def backups_restore():
    rel = request.form.get("rel") or ""
    try:
        result = backup_svc.restore_backup(rel, user_name=current_user_name())
        db.log_audit(
            current_user_name(),
            "استعادة بيانات",
            "نسخة احتياطية",
            details=f"استعادة {rel} | أمان: {result['safety'].get('id')}",
        )
        flash("تمت الاستعادة بنجاح (مع حفظة أمان تلقائية).", "ok")
    except Exception as exc:
        flash(f"تعذر الاستعادة: {exc}", "danger")
    return redirect(url_for("backups_home"))


@app.route("/admin/backups/upload", methods=["POST"])
@login_required
def backups_upload():
    f = request.files.get("backup_zip")
    if not f or not f.filename:
        flash("اختر ملف ZIP للحفظة.", "danger")
        return redirect(url_for("backups_home"))
    also_restore = request.form.get("also_restore") == "1"
    try:
        result = backup_svc.import_backup_zip(
            f,
            user_name=current_user_name(),
            also_restore=also_restore,
        )
        db.log_audit(
            current_user_name(),
            "رفع حفظة",
            "نسخة احتياطية",
            details=f"{result['imported'].get('label')} — استعادة={also_restore}",
        )
        flash("تم رفع الحفظة إلى السيرفر.", "ok")
        if also_restore:
            flash("تمت استعادة البيانات من الملف المرفوع.", "ok")
    except Exception as exc:
        flash(f"تعذر رفع الحفظة: {exc}", "danger")
    return redirect(url_for("backups_home"))


@app.route("/admin/backups/auto-now", methods=["POST"])
@login_required
def backups_auto_now():
    try:
        result = backup_svc.create_auto_backup(force=True, user_name=current_user_name())
        if result.get("created"):
            flash("تم إنشاء حفظة تلقائية وتجهيزها للجهاز الرئيسي.", "ok")
        else:
            flash(result.get("reason") or result.get("error") or "لم تُنشأ حفظة", "warning")
    except Exception as exc:
        flash(f"تعذر الحفظ التلقائي: {exc}", "danger")
    return redirect(url_for("backups_home"))


@app.route("/admin/backups/regen-token", methods=["POST"])
@login_required
def backups_regen_token():
    try:
        backup_svc.regenerate_sync_token()
        flash("تم توليد رمز مزامنة جديد — حدّث ملف الإعداد على الجهاز الرئيسي.", "ok")
    except Exception as exc:
        flash(f"تعذر توليد الرمز: {exc}", "danger")
    return redirect(url_for("backups_home"))


@app.route("/api/backups/latest")
def api_backups_latest():
    """سحب أحدث حفظة للجهاز الرئيسي (يتطلب رمز المزامنة)."""
    token = request.args.get("token") or request.headers.get("X-Backup-Token")
    if not backup_svc.token_matches(token):
        return {"ok": False, "error": "رمز غير صالح"}, 401
    try:
        backup_svc.create_auto_backup(force=False)
    except Exception:
        pass
    latest = backup_svc.latest_backup(purpose="auto") or backup_svc.latest_backup()
    if not latest:
        return {"ok": False, "error": "لا توجد حفظات بعد"}, 404
    zip_path = backup_svc.build_backup_zip(latest["_rel"])
    stamp = (latest.get("created_at") or "").replace(":", "").replace("T", "-")
    return send_file(
        zip_path,
        as_attachment=True,
        download_name=f"rekaz-auto-{stamp or 'latest'}.zip",
        mimetype="application/zip",
    )


@app.route("/api/backups/auto-run")
def api_backups_auto_run():
    """تشغيل دورة الحفظ التلقائي (من keep-alive أو الجهاز الرئيسي)."""
    token = request.args.get("token") or request.headers.get("X-Backup-Token")
    if not backup_svc.token_matches(token):
        return {"ok": False, "error": "رمز غير صالح"}, 401
    force = (request.args.get("force") or "").strip() in {"1", "true", "yes"}
    result = backup_svc.create_auto_backup(force=force)
    return {
        "ok": True,
        **{k: v for k, v in result.items() if k != "backup"},
        "backup_id": (result.get("backup") or {}).get("id"),
    }, 200


@app.route("/api/backups/sync-status")
def api_backups_sync_status():
    token = request.args.get("token") or request.headers.get("X-Backup-Token")
    if not backup_svc.token_matches(token):
        return {"ok": False, "error": "رمز غير صالح"}, 401
    return {"ok": True, **backup_svc.auto_status()}, 200


@app.route("/sop")
@login_required
def sop_page():
    conn = db.connect()
    rows = db.rows_to_dicts(conn.execute("SELECT * FROM sop ORDER BY id").fetchall())
    conn.close()
    return render_template("sop.html", rows=rows)


@app.route("/warehouses/balances")
@login_required
def warehouse_balances():
    q = (request.args.get("q") or "").strip().lower()
    conn = db.connect()
    items = db.rows_to_dicts(conn.execute("SELECT * FROM warehouse_items ORDER BY item_no").fetchall())
    conn.close()
    for item in items:
        detail = db.warehouse_balance_detail(item.get("item_no"))
        item.update(detail)
    if q:
        items = [
            r
            for r in items
            if q in (r.get("item_no") or "").lower()
            or q in (r.get("item_name") or "").lower()
            or q in (r.get("category") or "").lower()
        ]
    return render_template(
        "warehouse_balances.html",
        rows=items,
        q=q,
        section="warehouses",
        section_meta=SECTION_META["warehouses"],
        section_modules=modules_for_section("warehouses"),
    )


@app.route("/warehouses/balances/template.xlsx")
@login_required
def warehouse_items_template():
    data = warehouse_excel.build_items_template()
    return send_file(
        io.BytesIO(data),
        as_attachment=True,
        download_name="قالب_مواد_المستودع.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/warehouses/tx/template.xlsx")
@login_required
def warehouse_tx_template():
    data = warehouse_excel.build_tx_template()
    return send_file(
        io.BytesIO(data),
        as_attachment=True,
        download_name="قالب_حركات_المستودع.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/warehouses/balances/import", methods=["POST"])
@login_required
def warehouse_items_import():
    if not permissions.can("modules.write"):
        return permissions.deny_redirect()
    f = request.files.get("file")
    if not f or not f.filename:
        flash("اختر ملف Excel للمواد", "danger")
        return redirect(url_for("warehouse_balances"))
    try:
        result = warehouse_excel.import_items_from_excel(f)
        flash(
            f"استيراد المواد: جديد {result['ok']} | محدّث {result['updated']} | أرصدة افتتاحية {result['opening']}",
            "ok",
        )
        if result.get("errors"):
            flash(" / ".join(result["errors"][:5]), "danger")
        db.log_audit(current_user_name(), "استيراد Excel", "أصناف المستودع", details=str(result)[:240])
    except Exception as exc:
        flash(f"تعذر الاستيراد: {exc}", "danger")
    return redirect(url_for("warehouse_balances"))


@app.route("/warehouses/tx/import", methods=["POST"])
@login_required
def warehouse_tx_import():
    if not permissions.can("modules.write"):
        return permissions.deny_redirect()
    f = request.files.get("file")
    if not f or not f.filename:
        flash("اختر ملف Excel للحركات", "danger")
        return redirect(url_for("warehouse_balances"))
    try:
        result = warehouse_excel.import_tx_from_excel(f)
        flash(
            f"استيراد الحركات: {result['ok']} حركة | مربوطة ببلاغات: {result['linked']}",
            "ok",
        )
        if result.get("errors"):
            flash(" / ".join(result["errors"][:5]), "danger")
        db.log_audit(current_user_name(), "استيراد Excel", "معاملات المستودع", details=str(result)[:240])
    except Exception as exc:
        flash(f"تعذر الاستيراد: {exc}", "danger")
    return redirect(url_for("warehouse_balances"))


@app.route("/users/list", methods=["GET", "POST"])
@login_required
def users_list():
    conn = db.connect()
    if request.method == "POST":
        action = request.form.get("action")
        if action == "add":
            try:
                role = permissions.normalize_role(request.form.get("role") or "مدخل بيانات")
                conn.execute(
                    "INSERT INTO users(username, full_name, role, active, password, notes) VALUES (?,?,?,?,?,?)",
                    (
                        request.form.get("username"),
                        request.form.get("full_name"),
                        role,
                        1 if request.form.get("active") == "1" else 0,
                        request.form.get("password") or "1234",
                        request.form.get("notes"),
                    ),
                )
                conn.commit()
                db.log_audit(current_user_name(), "إضافة", "مستخدم", details=request.form.get("username"))
                flash("تم إضافة المستخدم", "ok")
            except Exception as exc:
                flash(f"تعذر الإضافة: {exc}", "danger")
        elif action == "update":
            uid = request.form.get("id")
            role = permissions.normalize_role(request.form.get("role") or "مدخل بيانات")
            password = (request.form.get("password") or "").strip()
            if password:
                conn.execute(
                    "UPDATE users SET full_name=?, role=?, active=?, password=?, notes=? WHERE id=?",
                    (
                        request.form.get("full_name"),
                        role,
                        1 if request.form.get("active") == "1" else 0,
                        password,
                        request.form.get("notes"),
                        uid,
                    ),
                )
            else:
                conn.execute(
                    "UPDATE users SET full_name=?, role=?, active=?, notes=? WHERE id=?",
                    (
                        request.form.get("full_name"),
                        role,
                        1 if request.form.get("active") == "1" else 0,
                        request.form.get("notes"),
                        uid,
                    ),
                )
            conn.commit()
            if str(session.get("user_id")) == str(uid):
                session["full_name"] = request.form.get("full_name")
                session["role"] = role
            db.log_audit(current_user_name(), "تعديل", "مستخدم", uid)
            flash("تم تحديث المستخدم", "ok")
        elif action == "delete":
            uid = request.form.get("id")
            target = conn.execute("SELECT * FROM users WHERE id=?", (uid,)).fetchone()
            if target and str(session.get("user_id")) == str(uid):
                flash("لا يمكن حذف حسابك الحالي", "danger")
            else:
                admins = conn.execute(
                    "SELECT COUNT(*) FROM users WHERE lower(role)='admin' AND active=1"
                ).fetchone()[0]
                if target and permissions.normalize_role(target["role"]) == "admin" and admins <= 1:
                    flash("لا يمكن حذف آخر مدير نظام نشط", "danger")
                else:
                    conn.execute("DELETE FROM users WHERE id=?", (uid,))
                    conn.commit()
                    db.log_audit(current_user_name(), "حذف", "مستخدم", uid)
                    flash("تم الحذف", "ok")
        elif action == "toggle":
            uid = request.form.get("id")
            target = conn.execute("SELECT * FROM users WHERE id=?", (uid,)).fetchone()
            if target and str(session.get("user_id")) == str(uid):
                flash("لا يمكن إيقاف حسابك الحالي", "danger")
            else:
                conn.execute(
                    "UPDATE users SET active = CASE WHEN active=1 THEN 0 ELSE 1 END WHERE id=?",
                    (uid,),
                )
                conn.commit()
                flash("تم تحديث الحالة", "ok")
    rows = db.rows_to_dicts(conn.execute("SELECT * FROM users ORDER BY id").fetchall())
    conn.close()
    for row in rows:
        row["role"] = permissions.normalize_role(row.get("role"))
        row["perm_count"] = len(permissions.perms_for_role(row["role"]))
    return render_template(
        "users.html",
        rows=rows,
        role_matrix=permissions.role_matrix(),
        perm_labels=permissions.PERM_LABELS,
    )


@app.route("/admin/audit-log/view")
@login_required
def audit_log_page():
    q = (request.args.get("q") or "").strip()
    conn = db.connect()
    if q:
        like = f"%{q}%"
        rows = db.rows_to_dicts(
            conn.execute(
                """
                SELECT * FROM audit_log
                WHERE user_name LIKE ? OR action LIKE ? OR entity LIKE ? OR details LIKE ?
                ORDER BY id DESC LIMIT 300
                """,
                (like, like, like, like),
            ).fetchall()
        )
    else:
        rows = db.rows_to_dicts(conn.execute("SELECT * FROM audit_log ORDER BY id DESC LIMIT 300").fetchall())
    conn.close()
    return render_template("audit_log.html", rows=rows, q=q)


@app.route("/search")
@login_required
def global_search():
    q = (request.args.get("q") or "").strip()
    results = []
    if q:
        like = f"%{q}%"
        conn = db.connect()
        for t in conn.execute(
            """
            SELECT id, ticket_no, district, fault_type, status FROM tickets
            WHERE ticket_no LIKE ? OR district LIKE ? OR fault_type LIKE ? OR station_no LIKE ?
            LIMIT 40
            """,
            (like, like, like, like),
        ).fetchall():
            results.append(
                {
                    "scope": "العمليات",
                    "tab": "بلاغات",
                    "label": t["ticket_no"],
                    "detail": f"{t['district'] or ''} — {t['fault_type'] or ''} — {t['status'] or ''}",
                    "url": url_for("ticket_view", ticket_id=t["id"]),
                }
            )
        for name, mod in MODULES.items():
            cols = [f[0] for f in mod["fields"] if f[2] in ("text", "ticket")][:3]
            if not cols:
                continue
            where = " OR ".join([f"CAST({c} AS TEXT) LIKE ?" for c in cols])
            try:
                rows = conn.execute(
                    f"SELECT id, {', '.join(cols)} FROM {mod['table']} WHERE {where} LIMIT 20",
                    tuple([like] * len(cols)),
                ).fetchall()
            except Exception:
                continue
            for r in rows:
                detail = " | ".join(str(r[c] or "") for c in cols)
                results.append(
                    {
                        "scope": SECTION_META.get(mod.get("section"), {}).get("title", ""),
                        "tab": mod["title"],
                        "label": detail[:80],
                        "detail": detail,
                        "url": url_for("module_edit", name=name, row_id=r["id"]),
                    }
                )
        conn.close()
    return render_template("search.html", q=q, results=results)


# ---------- Jump ----------
@app.route("/api/jump-destinations")
@login_required
def api_jump_destinations():
    from flask import jsonify

    return jsonify(permissions.filter_jump_items(review_engine.jump_destinations()))


@app.route("/export/tickets.xlsx")
@login_required
def export_tickets_excel():
    from openpyxl import Workbook

    conn = db.connect()
    rows = db.rows_to_dicts(conn.execute("SELECT * FROM tickets ORDER BY id").fetchall())
    conn.close()
    wb = Workbook()
    ws = wb.active
    ws.title = "Tracking"
    headers = [
        "رقم البلاغ",
        "تاريخ الاستلام",
        "الحي",
        "وقت الاستلام",
        "المندوب",
        "رقم المحطة",
        "رقم الفيدر",
        "الموقع",
        "نوع العطل",
        "تصنيف البلاغ",
        "الفرقة",
        "وقت التوجيه",
        "وقت الوصول",
        "حالة التنفيذ",
        "تاريخ التنفيذ",
        "قيمة البنود",
        "ملاحظات",
    ]
    ws.append(headers)
    for t in rows:
        ws.append(
            [
                t.get("ticket_no"),
                t.get("receive_date"),
                t.get("district"),
                t.get("receive_time"),
                t.get("agent"),
                t.get("station_no"),
                t.get("feeder_no"),
                t.get("location"),
                t.get("fault_type"),
                t.get("classification"),
                t.get("team"),
                t.get("dispatch_time"),
                t.get("arrival_time"),
                t.get("status"),
                t.get("execution_date"),
                t.get("items_value"),
                t.get("notes"),
            ]
        )
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name=f"rakaz-tickets-{datetime.now().strftime('%Y%m%d')}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def main():
    create_app()
    host = os.environ.get("HOST", "127.0.0.1")
    port = int(os.environ.get("PORT", "5070"))
    use_waitress = os.environ.get("USE_WAITRESS", "").lower() in {"1", "true", "yes"}
    # على Render وأمثاله نستخدم Waitress تلقائياً
    if use_waitress or os.environ.get("RENDER"):
        from waitress import serve

        print(f"تشغيل الإنتاج (Waitress) على http://{host}:{port}")
        serve(app, host=host, port=port, threads=8)
    else:
        print(f"تشغيل نظام متابعة الأعمال العام — مكتب خدمات خريص على http://{host}:{port}")
        app.run(host=host, port=port, debug=False)


if __name__ == "__main__":
    main()
