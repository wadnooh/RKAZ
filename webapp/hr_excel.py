"""استيراد وتصدير بيانات الموظفين من/إلى Excel — نموذج معتمد واستيراد ذكي."""

from __future__ import annotations

import io
from datetime import datetime, date, time
from openpyxl import Workbook, load_workbook

from webapp import db
from webapp import excel_brand as brand

HR_HEADERS = [
    "الرقم الوظيفي (*)",
    "الاسم الكامل (*)",
    "الإدارة",
    "القسم",
    "المسمى الوظيفي",
    "المهنة في الإقامة",
    "الجنسية",
    "الحالة",
    "رقم الجوال",
    "رقم الهوية / الإقامة",
    "تاريخ انتهاء الإقامة",
    "رقم رخصة القيادة",
    "تاريخ انتهاء رخصة القيادة",
    "تاريخ انتهاء التأمين الطبي",
    "تاريخ نهاية العقد",
    "تاريخ الالتحاق",
    "الراتب الأساسي",
    "بدل السكن",
    "بدلات أخرى",
    "اسم البنك",
    "رقم الآيبان",
    "مسؤول الطوارئ",
    "هاتف الطوارئ",
    "ملاحظات",
]

HR_FIELDS = [
    "emp_no",
    "full_name",
    "administration",
    "department",
    "job_title",
    "profession",
    "nationality",
    "status",
    "phone",
    "id_number",
    "id_expiry_date",
    "driving_license_no",
    "license_expiry_date",
    "insurance_expiry_date",
    "contract_end_date",
    "join_date",
    "basic_salary",
    "housing_allowance",
    "other_allowances",
    "bank_name",
    "iban",
    "emergency_contact_name",
    "emergency_contact_phone",
    "notes",
]

_HR_ALIASES = {
    "الرقم الوظيفي (*)": "emp_no",
    "الرقم الوظيفي": "emp_no",
    "رقم الموظف": "emp_no",
    "رقم وظيفي": "emp_no",
    "كود الموظف": "emp_no",
    "emp_no": "emp_no",
    "employee_no": "emp_no",
    "emp no": "emp_no",
    "id": "emp_no",

    "الاسم الكامل (*)": "full_name",
    "الاسم الكامل": "full_name",
    "اسم الموظف الكامل": "full_name",
    "الاسم الثلاثي": "full_name",
    "الاسم الرباعي": "full_name",
    "اسم الموظف": "full_name",
    "الاسم": "full_name",
    "الموظف": "full_name",
    "full_name": "full_name",
    "name": "full_name",
    "employee_name": "full_name",

    "الإدارة": "administration",
    "ادارة": "administration",
    "الإداره": "administration",
    "الادارة": "administration",
    "administration": "administration",
    "admin": "administration",

    "القسم": "department",
    "قسم": "department",
    "department": "department",
    "dept": "department",

    "المسمى الوظيفي": "job_title",
    "المسمى": "job_title",
    "الوظيفة": "job_title",
    "job_title": "job_title",
    "job title": "job_title",
    "title": "job_title",

    "المهنة في الإقامة": "profession",
    "المهنة": "profession",
    "مهنة الإقامة": "profession",
    "profession": "profession",

    "الجنسية": "nationality",
    "nationality": "nationality",

    "الحالة": "status",
    "الحالة الوظيفية": "status",
    "حالة الموظف": "status",
    "status": "status",

    "رقم الجوال": "phone",
    "الجوال": "phone",
    "الهاتف": "phone",
    "رقم الهاتف": "phone",
    "phone": "phone",
    "mobile": "phone",

    "رقم الهوية / الإقامة": "id_number",
    "الهوية الوطنية / الإقامة": "id_number",
    "رقم الهوية الوطنية": "id_number",
    "الهوية الوطنية": "id_number",
    "رقم الهوية": "id_number",
    "رقم الإقامة": "id_number",
    "الهوية": "id_number",
    "الإقامة": "id_number",
    "السجل المدني": "id_number",
    "id_number": "id_number",
    "iqama_no": "id_number",
    "national_id": "id_number",

    "تاريخ انتهاء الإقامة": "id_expiry_date",
    "انتهاء الإقامة": "id_expiry_date",
    "تاريخ انتهاء الهوية/الإقامة": "id_expiry_date",
    "تاريخ انتهاء الهوية": "id_expiry_date",
    "انتهاء الهوية": "id_expiry_date",
    "id_expiry_date": "id_expiry_date",
    "iqama_expiry": "id_expiry_date",

    "رقم رخصة القيادة": "driving_license_no",
    "رقم الرخصة": "driving_license_no",
    "الرخصة": "driving_license_no",
    "driving_license_no": "driving_license_no",
    "license_no": "driving_license_no",

    "تاريخ انتهاء رخصة القيادة": "license_expiry_date",
    "تاريخ انتهاء الرخصة": "license_expiry_date",
    "انتهاء الرخصة": "license_expiry_date",
    "license_expiry_date": "license_expiry_date",

    "تاريخ انتهاء التأمين الطبي": "insurance_expiry_date",
    "تاريخ انتهاء التأمين": "insurance_expiry_date",
    "انتهاء التأمين": "insurance_expiry_date",
    "insurance_expiry_date": "insurance_expiry_date",

    "تاريخ نهاية العقد": "contract_end_date",
    "تاريخ انتهاء العقد": "contract_end_date",
    "نهاية العقد": "contract_end_date",
    "انتهاء العقد": "contract_end_date",
    "contract_end_date": "contract_end_date",

    "تاريخ الالتحاق": "join_date",
    "تاريخ التعيين": "join_date",
    "تاريخ المباشرة": "join_date",
    "join_date": "join_date",

    "الراتب الأساسي": "basic_salary",
    "الأساسي": "basic_salary",
    "راتب أساسي": "basic_salary",
    "basic_salary": "basic_salary",
    "salary": "basic_salary",

    "بدل السكن": "housing_allowance",
    "سكن": "housing_allowance",
    "housing_allowance": "housing_allowance",

    "بدل النقل": "other_allowances",
    "بدل نقل": "other_allowances",
    "بدل المواصلات": "other_allowances",
    "بدلات أخرى": "other_allowances",
    "بدلات": "other_allowances",
    "other_allowances": "other_allowances",

    "اسم البنك": "bank_name",
    "البنك": "bank_name",
    "bank_name": "bank_name",
    "bank": "bank_name",

    "رقم الآيبان": "iban",
    "رقم الآيبان (IBAN)": "iban",
    "رقم الحساب / الآيبان": "iban",
    "رقم الحساب": "iban",
    "الآيبان": "iban",
    "ايبان": "iban",
    "iban": "iban",
    "IBAN": "iban",

    "مسؤول الطوارئ": "emergency_contact_name",
    "جهة الاتصال في الطوارئ": "emergency_contact_name",
    "جهة اتصال الطوارئ": "emergency_contact_name",
    "جهة الاتصال": "emergency_contact_name",
    "شخص للطوارئ": "emergency_contact_name",
    "اسم الطوارئ": "emergency_contact_name",
    "طوارئ": "emergency_contact_name",
    "emergency_contact_name": "emergency_contact_name",

    "هاتف الطوارئ": "emergency_contact_phone",
    "جوال الطوارئ": "emergency_contact_phone",
    "رقم طوارئ": "emergency_contact_phone",
    "emergency_contact_phone": "emergency_contact_phone",

    "ملاحظات": "notes",
    "الملاحظات": "notes",
    "notes": "notes",
}

_ADMINISTRATIONS = [
    "الإدارة العامة",
    "إدارة المشاريع",
    "الموارد البشرية",
    "الإدارة المالية",
]


def _clean_key(val) -> str:
    s = str(val or "").strip().lower()
    for ch in " \t\r\n*_()/-.,:;[]{}#@!~":
        s = s.replace(ch, "")
    return (
        s.replace("ـ", "")
        .replace("أ", "ا")
        .replace("إ", "ا")
        .replace("آ", "ا")
        .replace("ة", "ه")
        .replace("ى", "ي")
    )


_HR_CLEAN_MAP = {_clean_key(k): v for k, v in _HR_ALIASES.items()}


def _norm_text(val) -> str:
    return str(val or "").strip().lower().replace("ـ", "").replace("أ", "ا").replace("إ", "ا").replace("آ", "ا").replace("ة", "ه")


def _format_date(val) -> str:
    if val is None or val == "":
        return ""
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    if " " in s and len(s) >= 10:
        s = s.split(" ")[0]
    if "T" in s:
        s = s.split("T")[0]
    if "/" in s:
        parts = s.split("/")
        if len(parts) == 3:
            if len(parts[2]) == 4:
                return f"{parts[2]}-{parts[1].zfill(2)}-{parts[0].zfill(2)}"
            elif len(parts[0]) == 4:
                return f"{parts[0]}-{parts[1].zfill(2)}-{parts[2].zfill(2)}"
    return s[:10] if len(s) >= 10 and s[4] == "-" and s[7] == "-" else s


def _to_num(val) -> float | None:
    if val is None or val == "":
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace(",", "").replace("ر.س", "").replace("SAR", "").strip()
    try:
        return float(s)
    except (ValueError, TypeError):
        return None


def _clean_str(val) -> str:
    if val is None:
        return ""
    s = str(val).strip()
    return "" if s.lower() in ("none", "null", "nan") else s


def _normalize_administration(val: str) -> str:
    s = _clean_str(val)
    if not s:
        return "الموارد البشرية"
    norm = _norm_text(s)
    if "عام" in norm or "تنفيذ" in norm or "رئيس" in norm:
        return "الإدارة العامة"
    if "مشروع" in norm or "مشاريع" in norm:
        return "إدارة المشاريع"
    if "مال" in norm or "حساب" in norm:
        return "الإدارة المالية"
    if "موارد" in norm or "بشر" in norm or "توظيف" in norm or "شؤون" in norm:
        return "الموارد البشرية"
    return s


def build_hr_template() -> bytes:
    """توليد ملف Excel قالب رسمي معتمد لجلب وتعبئة بيانات الموظفين."""
    wb = Workbook()
    ws = wb.active
    ws.title = "نموذج بيانات الموظفين"
    ncol = len(HR_HEADERS)

    header_row = brand.apply_brand_header(
        ws,
        title="نموذج جلب وتحديث بيانات الموظفين",
        ncol=ncol,
        meta_lines=[
            "قالب معتمد لاستيراد وتحديث موظفي شركة ركاز المتقدمة للمقاولات",
            "الإدارات المعتمدة: الإدارة العامة · إدارة المشاريع · الموارد البشرية · الإدارة المالية",
            "صيغة التواريخ: YYYY-MM-DD (مثال: 2026-12-31)",
        ],
        summary_lines=[
            "إذا كان الرقم الوظيفي أو رقم الهوية مسجلاً مسبقاً يتم تحديث بياناته آلياً",
            "الحقول بعلامة (*) أساسية ومطلوبة لضمان صحة السجل",
        ],
    )

    brand.write_header_row(ws, HR_HEADERS, header_row)

    # صف توضيحي كمثال
    sample_row = [
        "1001",
        "عبدالله محمد السعيد",
        "الموارد البشرية",
        "شؤون الموظفين",
        "أخصائي موارد بشرية",
        "أخصائي موارد بشرية",
        "سعودي",
        "على رأس العمل",
        "0501234567",
        "1012345678",
        "2028-12-31",
        "DL-889900",
        "2029-05-15",
        "2027-08-20",
        "2027-12-31",
        "2023-01-01",
        8500,
        2125,
        500,
        "مصرف الراجحي",
        "SA0380000000000000000000",
        "سعد السعيد",
        "0507654321",
        "موظف مثبت على رأس العمل",
    ]

    r = header_row + 1
    for col_idx, val in enumerate(sample_row, start=1):
        cell = ws.cell(row=r, column=col_idx, value=val)
        if col_idx in (17, 18, 19) and isinstance(val, (int, float)):
            cell.number_format = "#,##0.00"

    brand.style_data_rows(ws, start_row=r, end_row=r, ncol=ncol)
    return brand.save_workbook_bytes(wb)


def import_employees_from_excel(file_storage, conn=None) -> dict:
    """
    قراءة ملف Excel لجلب بيانات الموظفين وتحديث القائم أو إدراج الجديد (UPSERT).
    """
    own = conn is None
    conn = conn or db.connect()
    try:
        db.ensure_schema(conn)

        wb = load_workbook(file_storage, data_only=True)
        ws = wb.active

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return {"ok": 0, "updated": 0, "total": 0, "errors": ["الملف فارغ لا يحتوي على بيانات"]}

        # 1. البحث عن صف الترويسة
        header_row_idx = -1
        col_mapping: dict[int, str] = {}

        for r_idx, row in enumerate(rows[:25]):
            if not row:
                continue
            mapping = {}
            for c_idx, cell_val in enumerate(row):
                val_str = str(cell_val or "").strip()
                if not val_str:
                    continue
                field = (
                    _HR_ALIASES.get(val_str)
                    or _HR_ALIASES.get(val_str.lower())
                    or _HR_CLEAN_MAP.get(_clean_key(val_str))
                )
                if field:
                    mapping[c_idx] = field
            if "full_name" in mapping.values() or "emp_no" in mapping.values() or "id_number" in mapping.values():
                header_row_idx = r_idx
                col_mapping = mapping
                break

        if header_row_idx == -1 or not col_mapping:
            return {
                "ok": 0,
                "updated": 0,
                "total": 0,
                "errors": ["تعذر التعرف على أعمدة بيانات الموظفين (يرجى التأكد من احتواء الملف على عمود الاسم أو الرقم الوظيفي أو الهوية)"],
            }

        new_count = 0
        updated_count = 0
        errors = []

        data_rows = rows[header_row_idx + 1 :]

        for line_no, row in enumerate(data_rows, start=header_row_idx + 2):
            if not row or all(c is None or str(c).strip() == "" for c in row):
                continue

            extracted: dict[str, any] = {}
            for col_idx, field_name in col_mapping.items():
                if col_idx < len(row):
                    raw = row[col_idx]
                    if field_name in ("basic_salary", "housing_allowance", "other_allowances"):
                        extracted[field_name] = _to_num(raw)
                    elif field_name in ("id_expiry_date", "license_expiry_date", "insurance_expiry_date", "contract_end_date", "join_date"):
                        extracted[field_name] = _format_date(raw)
                    elif field_name == "administration":
                        extracted[field_name] = _normalize_administration(raw)
                    else:
                        extracted[field_name] = _clean_str(raw)

            emp_no = extracted.get("emp_no") or ""
            full_name = extracted.get("full_name") or ""
            id_number = extracted.get("id_number") or ""

            # يجب توفر الاسم أو الرقم الوظيفي على الأقل
            if not full_name and not emp_no and not id_number:
                continue

            if not full_name and emp_no:
                full_name = f"موظف {emp_no}"

            # الحالة الافتراضية
            if not extracted.get("status"):
                extracted["status"] = "على رأس العمل"

            # فحص وجود الموظف سابقاً
            existing = None
            if emp_no:
                existing = conn.execute("SELECT id, emp_no FROM hr_employees WHERE emp_no=?", (emp_no,)).fetchone()
            if not existing and id_number:
                existing = conn.execute("SELECT id, emp_no FROM hr_employees WHERE id_number=?", (id_number,)).fetchone()

            try:
                if existing:
                    # تحديث السجل القائم
                    emp_id = existing["id"]
                    update_clauses = []
                    params = []
                    for f in HR_FIELDS:
                        if f in extracted and extracted[f] not in (None, ""):
                            update_clauses.append(f"{f} = ?")
                            params.append(extracted[f])
                    if update_clauses:
                        params.append(emp_id)
                        sql = f"UPDATE hr_employees SET {', '.join(update_clauses)} WHERE id = ?"
                        conn.execute(sql, params)
                        updated_count += 1
                else:
                    # إدراج موظف جديد
                    # إذا لم يكن هناك رقم وظيفي أصدر رقماً تلقائياً
                    if not emp_no:
                        last_id = conn.execute("SELECT MAX(id) FROM hr_employees").fetchone()[0] or 0
                        emp_no = f"EMP-{last_id + 1:04d}"
                        extracted["emp_no"] = emp_no

                    cols = []
                    placeholders = []
                    vals = []
                    for f in HR_FIELDS:
                        if f in extracted and extracted[f] is not None:
                            cols.append(f)
                            placeholders.append("?")
                            vals.append(extracted[f])

                    sql = f"INSERT INTO hr_employees ({', '.join(cols)}) VALUES ({', '.join(placeholders)})"
                    conn.execute(sql, vals)
                    new_count += 1

            except Exception as e:
                errors.append(f"سطر {line_no} ({full_name or emp_no}): {str(e)}")

        conn.commit()
        return {
            "ok": new_count,
            "updated": updated_count,
            "total": new_count + updated_count,
            "errors": errors,
        }

    finally:
        if own:
            conn.close()
